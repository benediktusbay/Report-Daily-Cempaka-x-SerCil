"""Closing-only Excel parsing and calculation; no database writes or file storage."""
import io
import re
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile

from openpyxl import load_workbook

SLOTS = {'so_bcc1': ('BCC1', 'Cempaka'), 'so_bc22': ('BC22', 'Cempaka'),
         'so_bc13': ('BC13', 'Serang'), 'so_bc66': ('BC66', 'Cilegon')}
DEPOT_CODES = {'BCC1': 'Cempaka', 'BC22': 'Cempaka', 'BC13': 'Serang', 'BC66': 'Cilegon'}
DEPOTS = ('Cempaka', 'Serang', 'Cilegon', 'Unassigned')
CATEGORIES = ('Device', 'Macbook', 'ACC')
TSH = {'Zefanya Septania Simorangkir', 'Rafhyski Alhasan', 'Ikmah Novtianingrum'}
# Audit reference only. Production targets always come from MonthlyTarget.
SEPTEMBER_2026_TARGET_REFERENCE = {
    'ALL': {'Device': 32_500_000_000, 'Macbook': 1_625_000_000,
            'ACC': 780_000_000, 'TOTAL': 34_905_000_000},
    'Cempaka': {'TOTAL': 32_220_000_000},
    'Serang': {'TOTAL': 819_264_432},
    'Cilegon': {'TOTAL': 1_865_735_568},
}
AUDIT_EXPLANATIONS = {
    'Depo dari Billing': ('Depo dari Billing', 'Dealer belum ada di Target Master. Depo diambil dari Cust Code pada Billing Baseline.', 'dealer'),
    'Depo master vs Billing berbeda': ('Depo master berbeda dari Billing', 'Closing memakai depo pada Target Master. Perbedaan dengan Cust Code Billing perlu diperiksa.', 'dealer'),
    'Depo fallback': ('Depo dari slot SO/DO', 'Dealer tidak ditemukan di Target Master maupun Billing Baseline. Depo mengikuti slot upload.', 'dealer'),
    'Unassigned': ('Depo belum diketahui', 'Dealer belum dapat dipetakan ke depo; nilainya tetap terlihat di ALL.', 'dealer'),
    'Cross-site': ('Transaksi lintas depo', 'Lokasi transaksi berbeda dari home depo dealer. Closing tetap mengikuti home depo.', 'dealer'),
    'Kategori dari deskripsi': ('Kategori dari deskripsi produk', 'Article Code tidak ditemukan pada Billing Baseline; kategori dikenali dari nama produk.', 'baris'),
    'Unmapped article': ('Produk belum dikenali', 'Produk ini belum masuk nilai Closing sampai kategorinya jelas.', 'baris'),
    'Duplicate row identik': ('Baris duplikat', 'Baris SO/DO yang sama muncul lebih dari sekali dan hanya dihitung sekali.', 'baris'),
    'Target vs acuan September berbeda': ('Target berbeda dari angka acuan', 'Target Master tetap dipakai untuk perhitungan. Tabel di bawah menunjukkan selisih terhadap acuan September.', 'pemeriksaan'),
    'Billing Baseline tidak sinkron': ('Tanggal Billing berbeda', 'Tanggal data Dashboard dan Billing Baseline tidak sama. Periksa apakah file yang diunggah sudah terbaru.', 'pemeriksaan'),
    'Dashboard vs Baseline Actual difference': ('Actual berbeda', 'Actual Dashboard berbeda dari total No Tax pada Billing Baseline untuk tiga salesman Closing.', 'pemeriksaan'),
    'Target Master periode belum tersedia': ('Target belum tersedia', 'Target Master belum berisi nilai untuk periode yang dipilih.', 'pemeriksaan'),
    'Actual Dashboard periode belum tersedia': ('Actual belum tersedia', 'Dashboard belum memiliki Billing untuk periode yang dipilih.', 'pemeriksaan'),
}

BILLING_COLUMNS = {
    'so': ('SO Number',), 'article_code': ('Article Code',),
    'description': ('Article Description',), 'brand': ('Brand Name', 'Brand'),
    'group': ('Item Group Desc', 'Item Group Description', 'Item Group'),
    'bp': ('Sold to Party Code',), 'dealer': ('Sold to Party Name',),
    'cust': ('Cust Code',), 'site': ('Site Code',), 'date': ('Billing Date',),
    'salesman': ('Salesman Name',),
    'amount': ('Total Nett Amount No Tax', 'Total Net Amount No Tax'),
    'qty': ('Quantity',),
}
SO_COLUMNS = {
    'so': ('Sales Order Number', 'SO Number'), 'do': ('DO Number', 'Delivery No'),
    'item': ('Item Number', 'SO Item No'), 'article_code': ('Article Code', 'Material'),
    'description': ('Article Description', 'Material Description', 'Description'),
    'bp': ('Soldto Code', 'Sold to Party Code', 'Sold to Code'),
    'dealer': ('Soldto Name', 'Sold to Party Name', 'Dealer'),
    'salesman': ('Salesman Name',), 'qty': ('Quantity', 'Qty'),
    'sales': ('TTL Sales Price',), 'discount': ('TTL Discount',),
    'date': ('Order Date',), 'site': ('Plant', 'Site Code'),
    'amount': ('Amount',), 'ppn': ('PPN',),
}
PPN_RATE = Decimal('1.11')
SALESMAN_ORDER = ('Rafhyski Alhasan', 'Zefanya Septania Simorangkir', 'Ikmah Novtianingrum')
DEPOT_GROUP = {'Cempaka': 'Cempaka', 'Serang': 'Sercil', 'Cilegon': 'Sercil'}
GROUP_LABEL = (('Cempaka', 'Cempaka'), ('Sercil', 'Serang & Cilegon'), ('Lainnya', 'Depo belum diketahui'))


def model_name(description):
    """Nama model ringkas untuk tampilan: 'iPhone 18 Pro Max', 'iPad 11 (A16)'."""
    text = re.sub(r'\s+', ' ', clean(description))
    m = re.match(r'(?i)^(iphone\s+\S+(?:\s+(?:pro|max|plus|mini))*)', text)
    if m:
        return re.sub(r'(?i)^iphone', 'iPhone', m.group(1))
    m = re.match(r'(?i)^(.*?)(?=[\s/]+(?:wi-?fi|cell(?:ular)?|\d+\s?(?:gb|tb))\b)', text)
    result = (m.group(1) if m and m.group(1) else text).strip()
    return result[:48] or 'Produk'

def clean(value):
    if value is None:
        return ''
    result = str(value).strip()
    return '' if result.casefold() in ('nan', 'none', 'nat', 'null') else result

def header_key(value):
    return re.sub(r'\s+', ' ', clean(value)).casefold()

def normalize_id(value):
    result = clean(value)
    if not result:
        return ''
    if re.fullmatch(r'\d+(?:\.\d+)?[eE][+-]?\d+', result):
        try:
            result = format(Decimal(result), 'f')
        except InvalidOperation:
            return ''
    if re.fullmatch(r'\d+\.0+', result):
        result = result.split('.')[0]
    return result.lstrip('0') or ('0' if result and set(result) == {'0'} else result)

def normalize_so(value):
    result = normalize_id(value)
    return result if re.fullmatch(r'\d+', result) else ''

def money(value, context, blank_zero=False):
    raw = clean(value)
    if not raw:
        if blank_zero:
            return Decimal(0)
        raise ValueError(f'{context}: nilai kosong.')
    raw = re.sub(r'(?i)rp|\s', '', raw)
    if ',' in raw and '.' in raw:
        raw = raw.replace('.', '').replace(',', '.') if raw.rfind(',') > raw.rfind('.') else raw.replace(',', '')
    elif ',' in raw:
        raw = raw.replace(',', '.') if len(raw.rsplit(',', 1)[-1]) in (1, 2) else raw.replace(',', '')
    elif re.fullmatch(r'-?\d{1,3}(?:\.\d{3})+', raw):
        raw = raw.replace('.', '')
    try:
        result = Decimal(raw)
    except InvalidOperation:
        raise ValueError(f'{context}: angka tidak valid ({value}).') from None
    if not result.is_finite():
        raise ValueError(f'{context}: angka tidak valid ({value}).')
    return result

def quantity_key(value):
    raw = clean(value)
    if not raw:
        return ''
    try:
        return format(Decimal(raw).normalize(), 'f')
    except InvalidOperation:
        return raw

def date_value(value):
    if isinstance(value, datetime):
        return value.date()
    if hasattr(value, 'year') and hasattr(value, 'month') and hasattr(value, 'day'):
        return value
    if isinstance(value, (int, float)):
        from openpyxl.utils.datetime import from_excel
        try:
            return from_excel(value).date()
        except (TypeError, ValueError, OverflowError):
            return None
    text = clean(value)
    if not text:
        return None
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None

def excel_records(stream, fields, required, label, sheet_name=None):
    stream.seek(0)
    try:
        wb = load_workbook(stream, read_only=True, data_only=True)
    except (BadZipFile, OSError, ValueError) as exc:
        raise ValueError(f'{label}: file Excel tidak dapat dibaca ({exc}).') from None
    try:
        if sheet_name and sheet_name not in wb.sheetnames:
            raise ValueError(f'{label}: sheet {sheet_name} tidak ditemukan.')
        sheet = wb[sheet_name] if sheet_name else wb.active
        rows = iter(sheet.values)
        index = None
        for _ in range(30):
            candidate = next(rows, None)
            if candidate is None:
                break
            found = {header_key(v): i for i, v in enumerate(candidate) if clean(v)}
            resolved = {name: next((found[header_key(alias)] for alias in choices if header_key(alias) in found), None)
                        for name, choices in fields.items()}
            if all(resolved[name] is not None for name in required):
                index = resolved
                break
        if index is None:
            raise ValueError(f'{label}: header tidak ditemukan. Perlu: {", ".join(fields[n][0] for n in required)}.')
        for line, row in enumerate(rows, 2):
            if not any(clean(v) for v in row):
                continue
            yield line, {name: row[i] if i is not None and i < len(row) else None for name, i in index.items()}
    finally:
        wb.close()

def depot_from_code(value, canonical_depo):
    raw = clean(value).upper()
    return DEPOT_CODES.get(raw, canonical_depo(value) if raw else 'Unassigned')

def apple_metadata(description, brand='', group='', classify=None):
    """Use Billing brand/group first; description inference is conservative and flagged."""
    text = clean(description)
    low = text.casefold()
    b = clean(brand).casefold()
    if b and b not in ('app', 'apple'):
        return 'Non-Apple', '', ''
    if b in ('app', 'apple') and classify:
        cat = classify(group)
        if cat in CATEGORIES:
            return 'Apple', cat, ''
    # Explicit non-Apple families override vague accessory words.
    if re.match(r'^(samsung|sam\s+galaxy|galaxy|xim\b|xiaomi|redmi|oppo|vivo|tecno|poco|infinix|realme|honor|huawei)\b', low):
        return 'Non-Apple', '', ''
    if re.match(r'^(iphone\b|ipad\b)', low):
        return 'Apple', 'Device', 'Kategori dari deskripsi'
    if re.match(r'^(macbook\b|mb\s+neo\b|mba\s+\d|mbp\s+\d|imac\b|mac\s+(mini|studio|pro)\b)', low):
        return 'Apple', 'Macbook', 'Kategori dari deskripsi'
    if re.match(r'^(app\b|apple\b|beats\b|airpods\b|earpods\b)', low):
        return 'Apple', 'ACC', 'Kategori dari deskripsi'
    return 'Unmapped', '', ''

def parse_billing(stream, month, classify, canonical_salesman, canonical_depo):
    billed_so = set()
    articles = {}
    cust_depot = {}
    dealer_names = {}
    bp_site = {}
    actual_rows = []
    rows = 0
    max_date = None
    for line, row in excel_records(stream, BILLING_COLUMNS,
                                   ('so', 'article_code', 'description', 'group', 'bp', 'cust', 'site',
                                    'date', 'salesman', 'amount'), 'Billing Baseline', 'Export'):
        so = normalize_so(row['so'])
        if not so:
            continue
        billed_so.add(so)  # ALL billed SO, before Apple/salesman/depot filters.
        rows += 1
        article = normalize_id(row['article_code'])
        brand = clean(row['brand'])
        group = clean(row['group'])
        desc = clean(row['description'])
        if article and (brand or group):
            articles[article] = (brand, group)
        bp = normalize_id(row['bp'])
        if bp and clean(row['dealer']):
            dealer_names[bp] = clean(row['dealer'])
        cust = depot_from_code(row['cust'], canonical_depo)
        site = clean(row['site']).upper()
        if bp and cust in DEPOTS[:3]:
            cust_depot[bp] = cust
        if bp and site:
            bp_site.setdefault(bp, set()).add(site)
        day = date_value(row['date'])
        if day and (max_date is None or day > max_date):
            max_date = day
        if not day or day.strftime('%Y-%m') != month:
            continue
        salesman = canonical_salesman(row['salesman'])
        if salesman not in TSH:
            continue
        identity, cat, flag = apple_metadata(desc, brand, group, classify)
        if identity != 'Apple' or cat not in CATEGORIES:
            continue
        actual_rows.append(dict(bp=bp, dealer=clean(row['dealer']), salesman=salesman,
                                category=cat, value=money(row['amount'], f'Billing baris {line}'),
                                site=site, cust=cust, so=so, flag=flag))
    return dict(month=month, billed_so=billed_so, articles=articles, cust_depot=cust_depot,
                dealer_names=dealer_names,
                bp_site=bp_site, actual_rows=actual_rows, max_date=max_date,
                rows=rows, so_count=len(billed_so))

def parse_so(stream, slot, month, canonical_salesman):
    if slot not in SLOTS:
        raise ValueError(f'Slot SO/DO tidak dikenal: {slot}.')
    label, source_depot = SLOTS[slot]
    records = []
    so_seen = set()
    for line, row in excel_records(stream, SO_COLUMNS,
                                   ('so', 'article_code', 'description', 'bp', 'salesman', 'sales', 'discount'),
                                   f'SO/DO {label}'):
        so = normalize_so(row['so'])
        if not so:
            continue
        day = date_value(row['date'])
        if day and day.strftime('%Y-%m') != month:
            continue
        owner = canonical_salesman(row['salesman'])
        if owner not in TSH:
            continue
        sales = money(row['sales'], f'SO/DO {label} baris {line} TTL Sales Price')
        discount = money(row['discount'], f'SO/DO {label} baris {line} TTL Discount', True)
        bp = normalize_id(row['bp'])
        article = normalize_id(row['article_code'])
        nett = sales + discount
        if clean(row.get('amount')):
            with_tax = money(row['amount'], f'SO/DO {label} baris {line} Amount', True)
        elif clean(row.get('ppn')):
            with_tax = nett + money(row['ppn'], f'SO/DO {label} baris {line} PPN', True)
        else:
            with_tax = (nett * PPN_RATE).quantize(Decimal(1))
        records.append(dict(nett_ppn=with_tax, order_date=day.isoformat() if day else '',
                            so=so, so_display=clean(row['so']), do=normalize_so(row['do']),
                            item=normalize_id(row['item']), article=article,
                            description=clean(row['description']), bp=bp,
                            dealer=clean(row['dealer']), salesman=owner, qty=clean(row['qty']),
                            qty_key=quantity_key(row['qty']),
                            sales=sales, discount=discount, nett=sales+discount,
                            source_slot=label, source_depot=source_depot,
                            site=clean(row['site']).upper()))
        so_seen.add(so)
    return dict(rows=records, row_count=len(records), so_count=len(so_seen))

def target_depot(target, canonical_depo):
    return depot_from_code(getattr(target, 'depo', ''), canonical_depo)

def calculate(baseline, slots, targets, billing_rows, classify, canonical_salesman, canonical_depo):
    """Return JSON-ready summary, details and audit. Caller provides two bounded DB reads."""
    master = {normalize_id(t.bp): t for t in targets if normalize_id(t.bp)}
    target_amount = defaultdict(Decimal)
    for t in targets:
        if canonical_salesman(t.salesman) not in TSH or not normalize_id(t.bp):
            continue
        depot = target_depot(t, canonical_depo)
        depot = depot if depot in DEPOTS else 'Unassigned'
        for cat, field in (('Device','device_target'), ('Macbook','macbook_target'), ('ACC','acc_target')):
            target_amount[(depot, cat)] += money(getattr(t, field) or 0, 'Target', True)
    warnings = Counter()
    examples = defaultdict(list)
    warning_seen = defaultdict(set)
    def dealer_name(bp, supplied=''):
        target = master.get(bp)
        return (clean(getattr(target, 'dealer', '')) if target else '') or baseline['dealer_names'].get(bp) or clean(supplied) or 'Dealer belum diketahui'
    def flag(name, detail='', unique=False, display=None):
        if unique and detail in warning_seen[name]:
            return
        if unique:
            warning_seen[name].add(detail)
        warnings[name] += 1
        if display and len(examples[name]) < 8:
            examples[name].append(display)
    if not any(value for value in target_amount.values()):
        flag('Target Master periode belum tersedia')
    if not billing_rows:
        flag('Actual Dashboard periode belum tersedia')
    def home(bp, site='', slot_depot='', supplied_dealer=''):
        dealer = dealer_name(bp, supplied_dealer)
        t = master.get(bp)
        if t:
            depot = target_depot(t, canonical_depo)
            depot = depot if depot in DEPOTS else 'Unassigned'
            cust = baseline['cust_depot'].get(bp)
            if cust and cust != depot:
                flag('Depo master vs Billing berbeda', bp, True,
                     f'{dealer}: master {depot}, Billing {cust}')
            source = 'Master'
        else:
            depot = baseline['cust_depot'].get(bp, '')
            if depot:
                source = 'Depo dari Billing'
                flag(source, bp, True, dealer)
            elif slot_depot:
                depot = slot_depot
                source = 'Depo fallback'
                flag(source, bp, True, dealer)
            else:
                depot = 'Unassigned'
                source = 'Unassigned'
                flag(source, bp, True, dealer)
        cross = bool(site and site in DEPOT_CODES and depot in DEPOTS[:3] and DEPOT_CODES[site] != depot)
        return depot, source, cross
    actual = defaultdict(Decimal)
    baseline_actual = defaultdict(Decimal)
    # Per salesman × depo × kategori (non-PPN dan PPN) + detail dealer per model.
    s_target = defaultdict(Decimal)
    s_actual, s_actual_ppn = defaultdict(Decimal), defaultdict(Decimal)
    s_pending, s_pending_ppn = defaultdict(Decimal), defaultdict(Decimal)
    s_nodo, s_nodo_ppn = defaultdict(Decimal), defaultdict(Decimal)
    dealers = {}

    def dealer_entry(bp, salesman, depot, supplied=''):
        entry = dealers.get(bp)
        if entry is None:
            entry = dealers[bp] = dict(bp=bp, dealer=dealer_name(bp, supplied), salesman=salesman,
                                       depot=depot, billing={}, do={}, no_do={}, owner=None, no_do_owner=None)
        return entry

    def add_model(bucket, cat, desc, qty, nett, with_tax, day='', so=''):
        name = model_name(desc)
        item = bucket.setdefault(f'{cat}|{name}', dict(category=cat, model=name, qty=0.0, nett=0.0,
                                                       ppn=0.0, dates=[], so=[]))
        try:
            item['qty'] += float(qty or 0)
        except (TypeError, ValueError):
            pass
        item['nett'] += float(nett)
        item['ppn'] += float(with_tax)
        if day and day not in item['dates']:
            item['dates'].append(day)
        if so and so not in item['so']:
            item['so'].append(so)

    for t in targets:
        salesman = canonical_salesman(t.salesman)
        if salesman not in TSH or not normalize_id(t.bp):
            continue
        depot = target_depot(t, canonical_depo)
        depot = depot if depot in DEPOTS else 'Unassigned'
        for cat, field in (('Device', 'device_target'), ('Macbook', 'macbook_target'), ('ACC', 'acc_target')):
            s_target[(salesman, depot, cat)] += money(getattr(t, field) or 0, 'Target', True)
    dashboard_date = None
    for row in billing_rows:
        day = row.billing_date
        if day and (dashboard_date is None or day > dashboard_date):
            dashboard_date = day
        salesman = canonical_salesman(row.salesman)
        if salesman not in TSH:
            continue
        identity, cat, _ = apple_metadata(row.article, '', row.item_group, classify)
        if identity != 'Apple' or cat not in CATEGORIES:
            continue
        bp = normalize_id(row.sold_to_code)
        depot, _, _ = home(bp)
        value = money(row.nett_amount or 0, 'Actual Dashboard', True)
        actual[(depot, cat)] += value
        with_tax = money(getattr(row, 'nett_amount_with_tax', None) or 0, 'Actual Dashboard PPN', True) \
            or (value * PPN_RATE).quantize(Decimal(1))
        key = (salesman, depot, cat)
        s_actual[key] += value
        s_actual_ppn[key] += with_tax
        entry = dealer_entry(bp, salesman, depot, getattr(row, 'sold_to_name', ''))
        add_model(entry['billing'], cat, row.article, getattr(row, 'quantity', 0) or 0, value, with_tax)
    for row in baseline['actual_rows']:
        depot, source, cross = home(row['bp'], row['site'], supplied_dealer=row['dealer'])
        baseline_actual[(depot, row['category'])] += row['value']
        if cross:
            flag('Cross-site', f"{row['bp']}: {row['site']}", True,
                 f"{dealer_name(row['bp'], row['dealer'])}: transaksi {row['site']}, home {depot}")
    processed = []
    seen = set()
    for slot in SLOTS:
        if slot not in slots:
            continue
        for row in slots[slot]['rows']:
            identity = (row['so'], row['item'], row['article'], row['qty_key'], row['sales'], row['discount'])
            if identity in seen:
                flag('Duplicate row identik', row['so'], display=dealer_name(row['bp'], row['dealer']))
                continue
            seen.add(identity)
            article = baseline['articles'].get(row['article'])
            if article:
                brand, group = article
                product, cat, category_flag = apple_metadata(row['description'], brand, group, classify)
            else:
                product, cat, category_flag = apple_metadata(row['description'], '', '', classify)
            if product == 'Non-Apple':
                continue
            depot, source, cross = home(row['bp'], row['source_slot'], row['source_depot'], row['dealer'])
            flags = []
            if source != 'Master':
                flags.append(source)
            if cross:
                flags.append(f"Cross-site: {row['source_slot']}")
                flag('Cross-site', f"{row['bp']}: {row['source_slot']}", True,
                     f"{dealer_name(row['bp'], row['dealer'])}: transaksi {row['source_slot']}, home {depot}")
            if category_flag:
                flags.append(category_flag)
                flag(category_flag, row['article'], display=f"{dealer_name(row['bp'], row['dealer'])}: {row['description']}")
            item = dict(row)
            item.update(depot=depot, category=cat, product=product,
                        flags=flags, billed=row['so'] in baseline['billed_so'],
                        status='DO tersedia' if row['do'] else 'No DO belum tersedia')
            if product == 'Unmapped' or not cat:
                flag('Unmapped article', row['article'], display=f"{dealer_name(row['bp'], row['dealer'])}: {row['description']}")
            processed.append(item)
    pending_amount = defaultdict(Decimal)
    detail = {'pending': [], 'billed': [], 'no_do': [], 'unmapped': []}
    unmapped_amount = Decimal(0)
    excluded_amount = Decimal(0)
    no_do_amount = Decimal(0)
    for row in processed:
        out = dict(depot=row['depot'], source_slot=row['source_slot'], salesman=row['salesman'],
                   dealer=row['dealer'], bp=row['bp'], so=row['so_display'], normalized_so=row['so'],
                   do=row['do'] or '-', article=row['article'], description=row['description'],
                   category=row['category'] or 'Unmapped', qty=row['qty'],
                   nett=float(row['nett']), nett_ppn=float(row['nett_ppn']),
                   order_date=row['order_date'], status=row['status'], flags=', '.join(row['flags']))
        key = (row['salesman'], row['depot'], row['category'])
        if row['billed']:
            detail['billed'].append(out)
            excluded_amount += row['nett']
        elif not row['do']:
            detail['no_do'].append(out)
            no_do_amount += row['nett']
            if row['category']:
                s_nodo[key] += row['nett']
                s_nodo_ppn[key] += row['nett_ppn']
                entry = dealer_entry(row['bp'], row['salesman'], row['depot'], row['dealer'])
                entry['no_do_owner'] = entry['no_do_owner'] or (row['salesman'], row['depot'])
                add_model(entry['no_do'], row['category'], row['description'], row['qty'],
                          row['nett'], row['nett_ppn'], row['order_date'], row['so_display'])
        elif row['product'] == 'Unmapped' or not row['category']:
            detail['unmapped'].append(out)
            unmapped_amount += row['nett']
        else:
            detail['pending'].append(out)
            pending_amount[(row['depot'], row['category'])] += row['nett']
            s_pending[key] += row['nett']
            s_pending_ppn[key] += row['nett_ppn']
            entry = dealer_entry(row['bp'], row['salesman'], row['depot'], row['dealer'])
            entry['owner'] = entry['owner'] or (row['salesman'], row['depot'])
            add_model(entry['do'], row['category'], row['description'], row['qty'],
                      row['nett'], row['nett_ppn'], row['order_date'], row['so_display'])
    summary = {}
    for depot in ('ALL', *DEPOTS):
        summary[depot] = {}
        depots = DEPOTS if depot == 'ALL' else (depot,)
        for category in (*CATEGORIES, 'TOTAL'):
            cats = CATEGORIES if category == 'TOTAL' else (category,)
            a = sum((actual[(d,c)] for d in depots for c in cats), Decimal(0))
            p = sum((pending_amount[(d,c)] for d in depots for c in cats), Decimal(0))
            t = sum((target_amount[(d,c)] for d in depots for c in cats), Decimal(0))
            e = a+p
            summary[depot][category] = dict(actual=float(a), pending=float(p), estimated=float(e),
                                             target=float(t), pct=float(e/t*100) if t else 0,
                                             gap=float(e-t))
    salesman_view = build_salesman_view(s_target, s_actual, s_actual_ppn, s_pending, s_pending_ppn,
                                        s_nodo, s_nodo_ppn)
    dealer_list = []
    for entry in dealers.values():
        if not (entry['do'] or entry['no_do']):
            continue  # detail dealer hanya untuk dealer yang punya SO/DO belum billing
        out = dict(entry)
        for bucket in ('billing', 'do', 'no_do'):
            out[bucket] = sorted(entry[bucket].values(),
                                 key=lambda m: (CATEGORIES.index(m['category']) if m['category'] in CATEGORIES else 9,
                                                -m['nett']))
        owner = entry['owner'] or entry['no_do_owner'] or (entry['salesman'], entry['depot'])
        out['owner'] = f'{owner[0]}|{owner[1]}'
        out['no_do_owner'] = f'{entry["no_do_owner"][0]}|{entry["no_do_owner"][1]}' if entry['no_do_owner'] else ''
        dealer_list.append(out)
    dashboard_total = sum(actual.values(), Decimal(0))
    baseline_total = sum(baseline_actual.values(), Decimal(0))
    difference = dashboard_total - baseline_total
    reconciliation = []
    for depot in DEPOTS:
        for cat in CATEGORIES:
            delta = actual[(depot, cat)] - baseline_actual[(depot, cat)]
            if delta:
                reconciliation.append(dict(depot=depot, category=cat, difference=float(delta)))
    target_reference_differences = []
    if baseline['month'] == '2026-09':
        for depot, cats in SEPTEMBER_2026_TARGET_REFERENCE.items():
            for cat, expected in cats.items():
                observed = summary[depot][cat]['target']
                if observed != expected:
                    target_reference_differences.append(dict(depot=depot, category=cat,
                                                             expected=expected, actual=observed,
                                                             difference=observed-expected))
        if target_reference_differences:
            flag('Target vs acuan September berbeda', str(len(target_reference_differences)))
    if baseline['max_date'] != dashboard_date:
        flag('Billing Baseline tidak sinkron', f"Dashboard {dashboard_date}; Baseline {baseline['max_date']}")
    if difference:
        flag('Dashboard vs Baseline Actual difference', str(difference))
    warning_items = []
    for name, count in warnings.items():
        title, explanation, unit = AUDIT_EXPLANATIONS.get(name, (name, '', 'kejadian'))
        warning_items.append(dict(title=title, explanation=explanation,
                                  count=count, unit=unit, dealers=examples.get(name, [])))
    audit = dict(source='Dashboard — nilai tanpa pajak',
                 dashboard_date=dashboard_date.isoformat() if dashboard_date else '-',
                 baseline_date=baseline['max_date'].isoformat() if baseline['max_date'] else '-',
                 dashboard_actual=float(dashboard_total), baseline_actual=float(baseline_total),
                 difference=float(difference), unmapped=float(unmapped_amount),
                 no_do=float(no_do_amount), no_do_so=len({r['normalized_so'] for r in detail['no_do']}),
                 excluded=float(excluded_amount), warnings=dict(warnings), examples=dict(examples),
                 warning_items=warning_items,
                 reconciliation=reconciliation,
                 target_reference_differences=target_reference_differences)
    return dict(summary=summary, detail=detail, audit=audit, salesman_view=salesman_view, dealers=dealer_list)


def build_salesman_view(s_target, s_actual, s_actual_ppn, s_pending, s_pending_ppn, s_nodo, s_nodo_ppn):
    """Baris per salesman × depo, subtotal per kelompok depo, dan Total Area untuk tiap kategori."""
    combos = {(k[0], k[1]) for source in (s_target, s_actual, s_pending, s_nodo) for k, v in source.items() if v}

    def group_of(depot):
        return DEPOT_GROUP.get(depot, 'Lainnya')

    def metrics(include):
        values = {}
        for category in (*CATEGORIES, 'TOTAL'):
            cats = CATEGORIES if category == 'TOTAL' else (category,)
            keys = [(s, d, c) for (s, d) in combos if include(s, d) for c in cats]
            total = {name: sum((src[k] for k in keys), Decimal(0)) for name, src in (
                ('target', s_target), ('actual', s_actual), ('actual_ppn', s_actual_ppn),
                ('pending', s_pending), ('pending_ppn', s_pending_ppn),
                ('no_do', s_nodo), ('no_do_ppn', s_nodo_ppn))}
            estimated = total['actual'] + total['pending']
            target = total['target']
            values[category] = {**{k: float(v) for k, v in total.items()},
                                'estimated': float(estimated),
                                'estimated_ppn': float(total['actual_ppn'] + total['pending_ppn']),
                                'actual_pct': float(total['actual'] / target * 100) if target else None,
                                'pct': float(estimated / target * 100) if target else None,
                                'gap': float(estimated - target)}
        return values

    def order(combo):
        salesman, depot = combo
        rank = SALESMAN_ORDER.index(salesman) if salesman in SALESMAN_ORDER else 9
        return (rank, DEPOTS.index(depot) if depot in DEPOTS else 9)

    rows, cards = [], []
    for group, label in GROUP_LABEL:
        members = sorted((c for c in combos if group_of(c[1]) == group), key=order)
        if not members:
            continue
        for salesman, depot in members:
            rows.append(dict(kind='salesman', key=f'{salesman}|{depot}', salesman=salesman,
                             short=salesman.split()[0], depot=depot, group=group,
                             values=metrics(lambda s, d, a=salesman, b=depot: (s, d) == (a, b))))
        group_values = metrics(lambda s, d, g=group: group_of(d) == g)
        if len(members) > 1:
            rows.append(dict(kind='group', key=group, label=f'Total {label}', group=group, values=group_values))
        cards.append(dict(key=group, label=label, values=group_values))
    total_values = metrics(lambda s, d: True)
    rows.append(dict(kind='total', key='ALL', label='Total Area', group='ALL', values=total_values))
    cards.insert(0, dict(key='ALL', label='Total Area', values=total_values))
    return dict(rows=rows, cards=cards)